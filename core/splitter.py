"""직원별 시트 분할과 엑셀/ZIP 파일 만들기 (관리팀 수작업 양식과 같은 모양)."""
from __future__ import annotations

import io
import re
import zipfile

import pandas as pd
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from .analytics import work_type_in, work_type_out
from .config import WorkRule

_INVALID_SHEET_CHARS = r"[\[\]\:\*\?\/\\]"
_INVALID_FILE_CHARS = r'[<>:"/\\|?*]'
_HEADER_FILL = PatternFill("solid", fgColor="DDEBF7")
_WARN_FILL = PatternFill("solid", fgColor="FCE4D6")
_LEAVE_FILL = PatternFill("solid", fgColor="E2EFDA")

SHEET_COLUMNS = ["요일", "발생일자", "발생시각", "이름", "근무형태", "모드", "인증", "근무지", "비고", "확인"]
RAW_COLUMNS = ["요일", "발생일자", "발생시각", "이름", "근무형태", "모드", "인증", "근무지", "비고"]
_PLACE_TYPES = ("출장", "외근", "교육")
# 날짜·요일만 남기고 비워 두는 예외 상태
BLANK_STATUSES = {"기록 없음", "출근 기록 없음", "퇴근 기록 없음", "휴일근무(기록 없음)"}
_BLANK_FILL = PatternFill("solid", fgColor="EDEDED")


def safe_sheet_name(name: str, used: set[str]) -> str:
    """엑셀 시트 이름 규칙(31자, 특수문자 금지, 중복 불가)에 맞게 바꾼다."""
    base = re.sub(_INVALID_SHEET_CHARS, "_", str(name)).strip("' ") or "Sheet"
    base = base[:31]
    candidate, i = base, 2
    while candidate.lower() in used:
        suffix = f"_{i}"
        candidate = base[: 31 - len(suffix)] + suffix
        i += 1
    used.add(candidate.lower())
    return candidate


def _one_line(text: str) -> str:
    return " ".join(str(text).split())


def work_type_label(kind: str, ts: pd.Timestamp, holiday: bool, filled: bool, rule: WorkRule) -> str:
    """근무형태 자동 표기. 수작업 시트처럼 특이한 날만 적는다 (rule.label_all이면 모든 행)."""
    if filled:  # 하이웍스로 채운 출장·교육 시각
        t = rule.start if kind == "출근" else rule.end
        return f"정상{kind}({t:%H:%M})"
    if holiday:
        return work_type_in(ts, True, rule) if kind == "출근" else work_type_out(ts, True, rule)
    label = work_type_in(ts, False, rule) if kind == "출근" else work_type_out(ts, False, rule)
    if label.startswith("연장") or rule.label_all:
        return label
    return ""


def _valid_id(v) -> bool:
    return v is not None and not (isinstance(v, float) and pd.isna(v))


def employee_rows(days: pd.DataFrame, records: pd.DataFrame, rule: WorkRule) -> pd.DataFrame:
    """한 직원의 일별 근태표 -> 날짜별 출근 1줄 + 퇴근 1줄 (기록 없는 날은 1줄)."""
    out = []
    for _, d in days.sort_values("일자").iterrows():
        types = [t for t in d["근태"].split(", ") if t]
        content = d["근태내용"]
        place = _one_line(content.split("\n")[0]) if any(t in _PLACE_TYPES or t == "휴일근무" for t in types) else ""
        mates = d["동행"].replace(", ", "\n")
        other_types = [t for t in types if t not in _PLACE_TYPES and t != "휴일근무"]
        leave_note = (f"{', '.join(other_types)}: {_one_line(content)}" if content else ", ".join(other_types)) \
            if other_types else ""
        note = mates or leave_note
        check = [] if d["상태"] in ("정상", "휴일근무") else [d["상태"]]
        if d["상태"] == "퇴근 기록 없음" and pd.notna(d["마지막태그"]) and d["마지막태그"] > d["출근"]:
            check = [f"퇴근 기록 없음(마지막 태그 {d['마지막태그']:%H:%M})"]
        check += [t for t in types if t not in check and t != "휴일근무"]
        check_text = ", ".join(check)
        day_text = d["일자"].strftime("%Y-%m-%d")
        ov_place, ov_note = _ov(d, "수정_근무지"), _ov(d, "수정_비고")

        if d["상태"] in BLANK_STATUSES:
            # 기록 없음·출근만·퇴근만 있는 날: 대개 휴가 또는 미연동이라 날짜·요일만 적고 나머지는 비운다
            out.append([d["요일"], day_text, "", "", "", "", "", ov_place or "", ov_note or "", ""])
            continue
        if pd.isna(d["출근"]) and pd.isna(d["퇴근"]):
            out.append([d["요일"], day_text, d["상태"], "", "", "", "",
                        place if ov_place is None else ov_place, note if ov_note is None else ov_note, check_text])
            continue
        for kind in ("출근", "퇴근"):
            ts = d[kind]
            if pd.isna(ts):
                continue
            row_id = d[f"{kind}행"]
            if _valid_id(row_id):
                rec = records.loc[int(row_id)]
                name, auth, raw_wt, raw_place, raw_note = (rec["이름"], rec["인증"], rec["근무형태"],
                                                           rec["근무지"], rec["비고"])
            else:  # 하이웍스로 채웠거나 웹에서 기입한 시각
                src = d[f"{kind}출처"] or "하이웍스"
                name, auth, raw_wt, raw_place, raw_note = (d["직원"], src, "", "", "")
            filled = d[f"{kind}출처"] == "하이웍스"
            wt = _ov(d, f"수정_{kind}근무형태")
            if wt is None and kind == "퇴근" and "연장근무" in types:
                m = re.search(r"~\s*(\d{1,2}:\d{2})", content)
                if m:
                    wt = f"연장({int(m.group(1).split(':')[0]):02d}:{m.group(1).split(':')[1]})"
            if wt is None:
                wt = raw_wt or work_type_label(kind, ts, d["휴일"], filled, rule)
            out.append([d["요일"], day_text, ts.strftime("%H:%M:%S"), name, wt, kind, auth,
                        (raw_place or place) if ov_place is None else ov_place,
                        (raw_note or note) if ov_note is None else ov_note, check_text])
    return pd.DataFrame(out, columns=SHEET_COLUMNS)


def _ov(d: pd.Series, col: str) -> str | None:
    """웹에서 수정한 값 (없으면 None)."""
    v = d.get(col)
    if v is None or (isinstance(v, float) and pd.isna(v)):
        return None
    return str(v)


def raw_rows(records: pd.DataFrame) -> pd.DataFrame:
    """캡스 원본과 같은 열 구성의 전체 기록."""
    return pd.DataFrame({
        "요일": records["요일"],
        "발생일자": records["출입일시"].dt.strftime("%Y-%m-%d"),
        "발생시각": records["출입일시"].dt.strftime("%H:%M:%S"),
        "이름": records["이름"].where(records["이름"] != "", records["직원"]),
        "근무형태": records["근무형태"],
        "모드": records["모드"],
        "인증": records["인증"],
        "근무지": records["근무지"],
        "비고": records["비고"],
    })


def _autofit(ws, min_row: int = 1) -> None:
    for col_cells in ws.iter_cols(min_row=min_row):
        lengths = [max(len(line) for line in str(c.value).split("\n")) for c in col_cells if c.value is not None]
        width = max(lengths, default=2)
        ws.column_dimensions[get_column_letter(col_cells[0].column)].width = min(max(width * 1.7, 6), 45)


def _style_header(ws, row: int) -> None:
    for cell in ws[row]:
        if cell.value is not None:
            cell.font = Font(bold=True)
            cell.fill = _HEADER_FILL
            cell.alignment = Alignment(horizontal="center")


def _write_employee_sheet(writer, sheet: str, title: str, rows: pd.DataFrame) -> None:
    # 수작업 양식: B2 제목, B4 헤더, B5부터 데이터
    rows.to_excel(writer, sheet_name=sheet, startrow=3, startcol=1, index=False)
    ws = writer.sheets[sheet]
    ws.cell(row=2, column=2, value=title).font = Font(bold=True, size=13)
    _style_header(ws, 4)
    check_col = 2 + SHEET_COLUMNS.index("확인")
    time_col = 2 + SHEET_COLUMNS.index("발생시각")
    for r in range(5, 5 + len(rows)):
        check = ws.cell(row=r, column=check_col).value
        time_val = str(ws.cell(row=r, column=time_col).value or "")
        if not time_val:  # 날짜·요일만 있는 예외 행: 연한 회색으로 표시만
            fill = _BLANK_FILL
        elif not time_val[0].isdigit():  # 연차·공가 등
            fill = _LEAVE_FILL
        elif check and any(k in check for k in ("지각", "조퇴")):
            fill = _WARN_FILL
        else:
            continue
        for c in range(2, 2 + len(SHEET_COLUMNS)):
            ws.cell(row=r, column=c).fill = fill
    ws.freeze_panes = "C5"
    _autofit(ws, min_row=4)


def build_workbook(
    records: pd.DataFrame,
    days: pd.DataFrame,
    employees: list[str],
    rule: WorkRule,
    summary: pd.DataFrame | None = None,
    label: str = "",
    include_raw: bool = True,
) -> bytes:
    """[N월 전체] + [요약] + 직원별 시트 엑셀(bytes)."""
    buf = io.BytesIO()
    used: set[str] = set()
    with pd.ExcelWriter(buf, engine="openpyxl") as writer:
        if include_raw:
            sheet = safe_sheet_name(f"{label} 전체".strip(), used)
            raw_rows(records).to_excel(writer, sheet_name=sheet, index=False)
            ws = writer.sheets[sheet]
            ws.freeze_panes = "A2"
            _style_header(ws, 1)
            _autofit(ws)
        if summary is not None and not summary.empty:
            sheet = safe_sheet_name("요약", used)
            summary.to_excel(writer, sheet_name=sheet, index=False)
            ws = writer.sheets[sheet]
            ws.freeze_panes = "B2"
            _style_header(ws, 1)
            _autofit(ws)
        for emp in employees:
            d = days[days["직원"] == emp]
            if d.empty:
                continue
            dept = d["부서"].iloc[0]
            title = f"{dept} {emp}".strip()
            _write_employee_sheet(writer, safe_sheet_name(emp, used), title, employee_rows(d, records, rule))
        if not used:  # 빈 통합문서 방지
            pd.DataFrame({"안내": ["선택한 조건에 맞는 기록이 없습니다."]}).to_excel(writer, sheet_name="안내", index=False)
    return buf.getvalue()


def build_zip(records: pd.DataFrame, days: pd.DataFrame, employees: list[str], rule: WorkRule,
              label: str = "") -> bytes:
    """직원 한 명당 엑셀 파일 하나씩 묶은 ZIP(bytes)."""
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as zf:
        used: set[str] = set()
        for emp in employees:
            if days[days["직원"] == emp].empty:
                continue
            fname = re.sub(_INVALID_FILE_CHARS, "_", f"{label} {emp}".strip())
            while fname.lower() in used:
                fname += "_"
            used.add(fname.lower())
            emp_records = records[records["직원"] == emp]
            data = build_workbook(emp_records, days, [emp], rule, include_raw=False, label=label)
            zf.writestr(f"{fname}.xlsx", data)
    return buf.getvalue()
