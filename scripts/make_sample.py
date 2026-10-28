"""실제 파일과 같은 형식의 가짜 샘플 2개를 만든다 (이름은 모두 가짜).

실행: python scripts/make_sample.py
  data/sample/9월 근태현황(샘플).xlsx   - 캡스 원본 시트 '9월 전체' + 직원별 정리 시트 1개
  data/sample/근태현황(샘플).xlsx       - '명단' + 날짜별 일일보고 시트
"""
from __future__ import annotations

import random
from datetime import date, datetime, timedelta
from pathlib import Path

from openpyxl import Workbook

WEEK = ["월", "화", "수", "목", "금", "토", "일"]
STAFF = [("품질팀", "가민준", "부장"), ("생산팀", "나서연", "차장"), ("생산팀", "다지호", "대리"),
         ("생산팀", "라수아", "사원"), ("연구소", "마도윤", "사원"), ("관리팀", "바하은", "사원")]
HOLIDAYS = {date(2026, 9, 24), date(2026, 9, 25), date(2026, 9, 26)}


def make_caps(path: Path, year: int = 2026, month: int = 9, seed: int = 7) -> None:
    rnd = random.Random(seed)
    rows = []
    d = date(year, month, 1)
    while d.month == month:
        off = d.weekday() >= 5 or d in HOLIDAYS
        base = datetime.combine(d, datetime.min.time())
        for dept, name, _ in STAFF:
            if (off and rnd.random() > 0.15) or (not off and rnd.random() < 0.05):
                continue
            t_in = base + timedelta(hours=7, minutes=rnd.randint(30, 75), seconds=rnd.randint(0, 59))
            t_out = base + timedelta(hours=17, minutes=rnd.randint(30, 170), seconds=rnd.randint(0, 59))
            rows.append((t_in - timedelta(seconds=30), name, "출입"))
            rows.append((t_in, name, "출근"))
            for _ in range(rnd.randint(0, 3)):
                rows.append((t_in + timedelta(minutes=rnd.randint(60, 480)), name, "출입"))
            if rnd.random() > 0.03:
                rows.append((t_out, name, "퇴근"))
        if d.weekday() == 4:  # 금요일 경비 설정
            rows.append((base + timedelta(hours=19), "경비", "세트"))
        d += timedelta(days=1)
    rows.sort()

    wb = Workbook()
    ws = wb.active
    ws.title = f"{month}월 전체"
    ws.append(["요일", "발생일자", "발생시각", "이름", "근무형태", "구분", "모드", "인증", "근무지", "비고"])
    for ts, name, mode in rows:
        ws.append([WEEK[ts.weekday()], ts.strftime("%Y-%m-%d"), ts.strftime("%H:%M:%S"), name, "", "일반",
                   mode, rnd.choice(["카드", "1:N 지문"]), "", ""])

    # 직원별로 정리해 둔 시트 (앱은 무시하고 원본 시트를 씀)
    p = wb.create_sheet(STAFF[1][1])
    p["B2"] = f"{STAFF[1][0]} {STAFF[1][1]}"
    p.append([])
    p.append([None, "요일", "발생일자", "발생시각", "이름", "근무형태", "구분", "모드", "인증", "근무지", "비고"])
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)


def make_hiworks(path: Path) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "명단"
    ws.append([])
    ws.append(["구분", "부서", "성명"])
    for i, (dept, name, _) in enumerate(STAFF, 1):
        ws.append([i, dept, name])

    events = {
        date(2026, 9, 9): [(STAFF[2], "연차", "개인사유")],
        date(2026, 9, 15): [(STAFF[0], "출장", "한국서부발전\n온도상승시험"), (STAFF[3], "반차", "개인사유\n13:00~17:30")],
        date(2026, 9, 22): [(STAFF[4], "외출", "개인사유\n08:30~10:00"), (STAFF[5], "교육", "상공회의소\n노무교육")],
    }
    types = ["연차", "반차", "반반차", "외출", "공가", "출장", "교육"]
    for day, items in sorted(events.items(), reverse=True):
        s = wb.create_sheet(f"{day:%y.%m.%d}({WEEK[day.weekday()]})")
        s.append([])
        s.append([None, "구분", "부서", "직급", "성명", "근태현황"] + [None] * 6 + ["근무지\n(목적지/사유/시간)", "비고"])
        s.append([None] * 5 + types)
        for i, ((dept, name, rank), kind, content) in enumerate(items, 1):
            flags = ["o" if t == kind else None for t in types]
            s.append([None, i, dept, rank, name] + flags + [content, None])
        s.append([])
        s.append([None, "휴가일정 관련"])
    wb.save(path)


if __name__ == "__main__":
    out = Path(__file__).resolve().parent.parent / "data" / "sample"
    make_caps(out / "9월 근태현황(샘플).xlsx")
    make_hiworks(out / "근태현황(샘플).xlsx")
    print(f"샘플 파일 생성: {out}")
