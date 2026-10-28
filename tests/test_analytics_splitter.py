"""근태 판정·근무형태·직원별 시트 테스트 (이름은 가짜)."""
import io
import zipfile
from datetime import date

import pandas as pd
from openpyxl import load_workbook

from core.analytics import build_days, employee_summary, hourly_by_weekday, work_type_in, work_type_out
from core.config import HOLIDAYS_2026, WorkRule
from core.hiworks import LEAVE_COLUMNS
from core.loader import apply_header, guess_mapping, normalize
from core.splitter import build_workbook, build_zip, employee_rows, safe_sheet_name

RULE = WorkRule()  # 08:30 ~ 17:30
HEADER = ["요일", "발생일자", "발생시각", "이름", "근무형태", "구분", "모드", "인증", "근무지", "비고"]


def _records():
    rows = [
        # 가나다: 9/1(화) ~ 9/7(월)
        ("2026-09-01", "08:20:00", "가나다", "출근"),
        ("2026-09-01", "12:00:00", "가나다", "해제"),  # 경비 해제: 판정 제외
        ("2026-09-01", "17:35:00", "가나다", "퇴근"),
        ("2026-09-02", "08:45:00", "가나다", "출근"),  # 지각
        ("2026-09-02", "19:10:00", "가나다", "퇴근"),  # 연장(19:00)
        # 9/3 기록 없음
        ("2026-09-04", "09:30:00", "가나다", "출근"),  # 반반차라 지각 아님
        ("2026-09-04", "17:31:00", "가나다", "퇴근"),
        ("2026-09-05", "08:10:00", "가나다", "출근"),  # 토요일
        ("2026-09-05", "15:20:00", "가나다", "퇴근"),
        ("2026-09-07", "08:00:00", "가나다", "출입"),
        ("2026-09-07", "08:01:00", "가나다", "출근"),
        ("2026-09-07", "17:40:00", "가나다", "퇴근"),
        # 라마바: 9/1 ~ 9/3
        ("2026-09-01", "08:00:00", "라마바", "출근"),
        ("2026-09-01", "16:00:00", "라마바", "퇴근"),  # 조퇴
        # 9/2 연차
        ("2026-09-03", "08:10:00", "라마바", "출근"),  # 퇴근 기록 없음
    ]
    raw = pd.DataFrame([HEADER] + [["", d, t, n, "", "일반", m, "카드", "", ""] for d, t, n, m in rows])
    recs, _ = normalize(apply_header(raw, 0), guess_mapping(HEADER))
    recs["부서"] = recs["이름"].map({"가나다": "생산팀", "라마바": "품질팀"})
    return recs


def _leaves():
    rows = [
        (date(2026, 9, 4), "가나다", "생산팀", "사원", "반반차", "개인사유\n08:30~10:30", "", ""),
        (date(2026, 9, 2), "라마바", "품질팀", "과장", "연차", "개인사유", "", ""),
    ]
    return pd.DataFrame(rows, columns=LEAVE_COLUMNS)


def _days():
    return build_days(_records(), _leaves(), RULE, HOLIDAYS_2026)


def _status(days, emp, d):
    return days[(days["직원"] == emp) & (days["일자"] == d)].iloc[0]


def test_build_days_status():
    days = _days()
    assert _status(days, "가나다", date(2026, 9, 1))["상태"] == "정상"
    assert _status(days, "가나다", date(2026, 9, 2))["상태"] == "지각"
    assert _status(days, "가나다", date(2026, 9, 3))["상태"] == "기록 없음"
    assert _status(days, "가나다", date(2026, 9, 4))["상태"] == "정상"
    assert _status(days, "가나다", date(2026, 9, 5))["상태"] == "휴일근무"
    assert not (days["일자"] == date(2026, 9, 6)).any()  # 일요일, 기록 없음 -> 행 없음
    assert _status(days, "라마바", date(2026, 9, 1))["상태"] == "조퇴"
    assert _status(days, "라마바", date(2026, 9, 2))["상태"] == "연차"
    assert _status(days, "라마바", date(2026, 9, 3))["상태"] == "퇴근 기록 없음"
    # 출입 태그보다 '출근' 모드 우선
    assert _status(days, "가나다", date(2026, 9, 7))["출근"] == pd.Timestamp("2026-09-07 08:01")


def test_work_type_labels():
    ts = pd.Timestamp
    assert work_type_in(ts("2026-09-01 08:20"), False, RULE) == "정상(08:30)"
    assert work_type_in(ts("2026-09-05 08:10"), True, RULE) == "휴일(08:30)"
    assert work_type_out(ts("2026-09-01 18:08"), False, RULE) == "정상(17:30)"
    assert work_type_out(ts("2026-09-01 20:11"), False, RULE) == "연장(20:00)"
    assert work_type_out(ts("2026-09-05 15:20"), True, RULE) == "휴일(15:00)"


def test_employee_summary():
    s = employee_summary(_days()).set_index("직원")
    assert s.loc["가나다", "지각"] == 1
    assert s.loc["가나다", "기록 없음"] == 1
    assert s.loc["가나다", "휴일근무"] == 1
    assert s.loc["가나다", "휴가사용(일)"] == 0.25
    assert s.loc["라마바", "휴가사용(일)"] == 1.0
    assert s.loc["라마바", "조퇴"] == 1


def test_heatmap_excludes_alarm():
    assert hourly_by_weekday(_records()).loc["화", "12시"] == 0


def test_employee_rows_format():
    recs, days = _records(), _days()
    rows = employee_rows(days[days["직원"] == "라마바"], recs, RULE)
    assert list(rows["발생시각"]) == ["08:00:00", "16:00:00", "연차", ""]
    assert list(rows["모드"]) == ["출근", "퇴근", "", ""]
    # 출근만 있는 날(9/3)은 날짜·요일만 남기고 나머지는 빈칸
    last = rows.iloc[-1]
    assert (last["요일"], last["발생일자"]) == ("목", "2026-09-03")
    assert all(v == "" for v in last.drop(["요일", "발생일자"]))
    assert rows["확인"].iloc[0] == "조퇴"

    rows = employee_rows(days[days["직원"] == "가나다"], recs, RULE)
    sep2 = rows[rows["발생일자"] == "2026-09-02"]
    assert list(sep2["근무형태"]) == ["", "연장(19:00)"]  # 평범한 출근은 비워 둠
    sep5 = rows[rows["발생일자"] == "2026-09-05"]
    assert list(sep5["근무형태"]) == ["휴일(08:30)", "휴일(15:00)"]
    sep4 = rows[rows["발생일자"] == "2026-09-04"]
    assert sep4["비고"].iloc[0].startswith("반반차")


def test_safe_sheet_name():
    used: set[str] = set()
    assert safe_sheet_name("a/b:c", used) == "a_b_c"
    assert safe_sheet_name("A_B_C", used) == "A_B_C_2"
    assert len(safe_sheet_name("가" * 40, used)) == 31


def test_build_workbook_layout():
    recs, days = _records(), _days()
    data = build_workbook(recs, days, ["가나다", "라마바"], RULE, summary=employee_summary(days), label="9월")
    wb = load_workbook(io.BytesIO(data))
    assert wb.sheetnames == ["9월 전체", "요약", "가나다", "라마바"]
    ws = wb["가나다"]
    assert ws["B2"].value == "생산팀 가나다"
    assert ws["B4"].value == "요일"
    assert ws["D5"].value == "08:20:00"
    assert wb["9월 전체"].max_row == len(recs) + 1


def test_build_zip():
    recs, days = _records(), _days()
    names = zipfile.ZipFile(io.BytesIO(build_zip(recs, days, ["가나다", "라마바"], RULE, label="9월"))).namelist()
    assert sorted(names) == ["9월 가나다.xlsx", "9월 라마바.xlsx"]


def _mini(rows):
    raw = pd.DataFrame([HEADER] + [["", d, t, n, "", "일반", m, "카드", "", ""] for d, t, n, m in rows])
    recs, _ = normalize(apply_header(raw, 0), guess_mapping(HEADER))
    recs["부서"] = "생산팀"
    return recs


def test_only_out_tag_is_out_and_travel_fill_and_companions():
    recs = _mini([
        ("2026-09-07", "17:44:33", "가나다", "퇴근"),   # 교육 다녀와서 퇴근만 찍음
        ("2026-09-08", "08:20:00", "가나다", "출근"),   # 출장: 오전에 나감
        ("2026-09-08", "10:25:00", "가나다", "출입"),
        ("2026-09-08", "08:10:00", "라마바", "출근"),
        ("2026-09-08", "17:40:00", "라마바", "퇴근"),
        ("2026-09-09", "08:00:00", "라마바", "출근"),
        ("2026-09-09", "17:41:00", "라마바", "출입"),  # 퇴근 모드 없음
    ])
    leaves = pd.DataFrame([
        (date(2026, 9, 7), "가나다", "생산팀", "", "교육", "서산 운전학원\n지게차", "", ""),
        (date(2026, 9, 8), "가나다", "생산팀", "", "출장", "원동취수장\n교체작업", "", ""),
        (date(2026, 9, 8), "다라마", "생산팀", "", "출장", "원동취수장\n교체작업", "", ""),
    ], columns=LEAVE_COLUMNS)
    days = build_days(recs, leaves, RULE, HOLIDAYS_2026)

    d7 = _status(days, "가나다", date(2026, 9, 7))
    assert d7["출근출처"] == "하이웍스" and d7["출근"] == pd.Timestamp("2026-09-07 08:30")
    assert d7["퇴근"] == pd.Timestamp("2026-09-07 17:44:33")
    d8 = _status(days, "가나다", date(2026, 9, 8))
    assert d8["퇴근출처"] == "하이웍스" and d8["퇴근"] == pd.Timestamp("2026-09-08 17:30")
    assert d8["동행"] == "다라마"
    assert _status(days, "라마바", date(2026, 9, 9))["상태"] == "퇴근 기록 없음"

    rows = employee_rows(days[days["직원"] == "가나다"], recs, RULE)
    assert list(rows["근무형태"]) == ["정상출근(08:30)", "", "", "정상퇴근(17:30)"]
    assert list(rows["인증"]) == ["하이웍스", "카드", "카드", "하이웍스"]
    assert rows["근무지"].iloc[0] == "서산 운전학원"
    assert rows["비고"].iloc[2] == "다라마"
    r = employee_rows(days[days["직원"] == "라마바"], recs, RULE)
    assert list(r.iloc[-1]) == ["수", "2026-09-09"] + [""] * 8


def test_no_record_day_is_date_only():
    recs, days = _records(), _days()
    rows = employee_rows(days[days["직원"] == "가나다"], recs, RULE)
    sep3 = rows[rows["발생일자"] == "2026-09-03"]
    assert len(sep3) == 1
    assert all(v == "" for v in sep3.iloc[0].drop(["요일", "발생일자"]))


def test_dashboard_aggregates():
    from core.analytics import attention, day_snapshot, leave_status, missing_rows, work_hours_report
    days = _days()
    snap = day_snapshot(days, date(2026, 9, 2))
    assert snap == {"출근": 1, "지각": 1, "미출근": 0, "휴가": 1}  # 가나다 지각, 라마바 연차

    miss = missing_rows(days)
    assert set(zip(miss["직원"], miss["표시"])) == {("가나다", "미출근"), ("라마바", "퇴근 누락")}
    assert miss["일자"].iloc[0] >= miss["일자"].iloc[-1]  # 최근 날짜 먼저

    rep = work_hours_report(days, RULE).set_index("직원")
    # 가나다 평일 9/1·2·3·4·7 = 5일 x 8h
    assert rep.loc["가나다", "소정근로(h)"] == 40
    # 9/2 19:10 퇴근 -> 17:30 이후 100분 -> 30분 단위 90분
    assert rep.loc["가나다", "연장근로(h)"] == 1.5
    # 9/5(토) 08:10~15:20 = 7h10m - 휴게 30분 = 6h40m -> 30분 단위 6.5h
    assert rep.loc["가나다", "휴일근로(h)"] == 6.5

    ls = leave_status(_leaves(), ["가나다", "라마바"], date(2026, 9, 30)).set_index("이름")
    assert ls.loc["가나다", "연차사용"] == 0.25 and ls.loc["라마바", "연차사용"] == 1.0

    groups = dict(attention(days, work_hours_report(days, RULE), ["경비"]))
    assert groups["휴일 출근 직원"] == ["가나다"]
    assert groups["명단에 없는 출입자"] == ["경비"]


def test_hm_format():
    from core.ui import hm
    assert hm(129 + 13 / 60) == "129h 13m"
    assert hm(2.5) == "2h 30m"
    assert hm(0) == "0h"


def test_actual_hours_ignore_early_arrival():
    from core.analytics import work_hours_report
    recs = _mini([("2026-09-01", "07:30:00", "가나다", "출근"), ("2026-09-01", "17:30:00", "가나다", "퇴근")])
    rep = work_hours_report(build_days(recs, None, RULE, HOLIDAYS_2026), RULE).set_index("직원")
    assert rep.loc["가나다", "실제근로(h)"] == 8.0  # 08:30~17:30 - 휴게 1h
